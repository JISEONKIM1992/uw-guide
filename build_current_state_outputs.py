from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


BASE = Path(r"C:\Users\user\Desktop")
WORKDIR = BASE / "uw guide"
OUTPUTS = WORKDIR / "outputs"

GUIDELINE_PATH = BASE / "기획agent_가이드라인v2.xlsx"
RD_PATH = BASE / "RD판.xlsx"

GUIDELINE_OUT = OUTPUTS / "기획agent_가이드라인v2_지원Agent정리.xlsx"
RD_OUT = OUTPUTS / "RD판_지원Agent_v2_요약시트추가.xlsx"

YELLOW_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")
ORANGE_FILL = PatternFill(fill_type="solid", fgColor="FCE4D6")
GREEN_FILL = PatternFill(fill_type="solid", fgColor="E2F0D9")
BLUE_FILL = PatternFill(fill_type="solid", fgColor="D9EAF7")
TITLE_FILL = PatternFill(fill_type="solid", fgColor="1F3A5F")
WHITE_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(color="FFFFFF", bold=True, size=14)
HEADER_FONT = Font(bold=True, color="1F3A5F")
BODY_FONT = Font(color="2F3B52")
THIN = Side(style="thin", color="B7C9DD")
BODY_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
TOP = Alignment(vertical="top", wrap_text=True)


def norm(value):
    return "" if value is None else str(value).strip()


def load_products():
    wb = load_workbook(GUIDELINE_PATH, data_only=False)
    ws = wb["Rule Master"]
    products = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        product_name = norm(row[0])
        product_code = norm(row[1])
        rule_id = norm(row[3])
        status = norm(row[4])
        rider_name = norm(row[5])
        insurance_code = norm(row[6])
        as_is_general = norm(row[7])
        to_be_general = norm(row[8])
        as_is_simple = norm(row[9])
        to_be_simple = norm(row[10])
        memo = norm(row[14])
        products.append(
            {
                "productName": product_name,
                "productCode": product_code,
                "ruleId": rule_id,
                "status": status,
                "riderName": rider_name,
                "insuranceCode": insurance_code,
                "generalAsIs": as_is_general,
                "generalToBe": to_be_general,
                "simpleAsIs": as_is_simple,
                "simpleToBe": to_be_simple,
                "memo": memo,
            }
        )
    return products


def label_name(product):
    return product.get("riderName") or product.get("productName") or "특약"


def format_limit_pair(product):
    g_as = product.get("generalAsIs") or "-"
    s_as = product.get("simpleAsIs") or "-"
    g_to = product.get("generalToBe") or product.get("generalAsIs") or "-"
    s_to = product.get("simpleToBe") or product.get("simpleAsIs") or "-"
    changed = (
        (product.get("generalToBe") and product["generalToBe"] != product.get("generalAsIs"))
        or (product.get("simpleToBe") and product["simpleToBe"] != product.get("simpleAsIs"))
    )
    if not changed:
        return None
    return f"{g_as}~{s_as} > {g_to}~{s_to}"


def classify_limit_change(product):
    changed = (
        (product.get("generalToBe") and product["generalToBe"] != product.get("generalAsIs"))
        or (product.get("simpleToBe") and product["simpleToBe"] != product.get("simpleAsIs"))
    )
    if not changed:
        return None
    aligned = bool(product.get("generalToBe")) and bool(product.get("simpleToBe")) and norm(product.get("generalToBe")) == norm(product.get("simpleToBe"))
    return {
        "status": "자동반영" if aligned else "룰검토필요",
        "reason": "일반/건강과 간편 한도가 동일" if aligned else "일반/건강과 간편 한도가 상이",
    }


def review_note(product):
    parts = []
    if product.get("generalAsIs") or product.get("simpleAsIs"):
        parts.append(f"as-is {product.get('generalAsIs') or '-'} / {product.get('simpleAsIs') or '-'}")
    if product.get("generalToBe") or product.get("simpleToBe"):
        parts.append(f"to-be {product.get('generalToBe') or '-'} / {product.get('simpleToBe') or '-'}")
    return " | ".join(parts)


def find_row_for_product(sheet, product):
    target_insurance = norm(product.get("insuranceCode"))
    target_name = norm(label_name(product))
    for row in range(5, sheet.max_row + 1):
        insurance = norm(sheet.cell(row, 6).value)
        name = norm(sheet.cell(row, 4).value)
        if target_insurance and insurance == target_insurance:
            return row
        if not target_insurance and target_name and name == target_name:
            return row
        if target_insurance and not insurance and target_name and name == target_name:
            return row
    return None


def highlight_changed(sheet, refs, fill):
    for ref in refs:
        cell = sheet[ref]
        cell.fill = fill
        cell.font = Font(bold=True, color="7F6000")
        cell.alignment = CENTER
        cell.border = BODY_BORDER


def style_table(sheet, title_cells, header_row, data_start, data_end, col_widths, title_fill=TITLE_FILL):
    for cell_range in title_cells:
        for row in sheet[cell_range]:
            for cell in row:
                cell.fill = title_fill
                cell.font = TITLE_FONT
                cell.alignment = CENTER
                cell.border = BODY_BORDER
    for cell in sheet[header_row]:
        cell.fill = BLUE_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BODY_BORDER
    for row in sheet.iter_rows(min_row=data_start, max_row=data_end):
        for cell in row:
            cell.border = BODY_BORDER
            cell.alignment = TOP
            cell.font = BODY_FONT
    for col, width in col_widths.items():
        sheet.column_dimensions[col].width = width


def build_rd_workbook(products):
    wb = load_workbook(RD_PATH)
    ws = wb["Sheet1"]
    changed_refs = []

    for product in products:
        row = find_row_for_product(ws, product)
        if row is None:
            continue
        current_product_code = norm(ws.cell(row, 2).value)
        current_insurance_code = norm(ws.cell(row, 6).value)
        new_limit = format_limit_pair(product)
        if product.get("productCode") and current_product_code != product["productCode"]:
            ws.cell(row, 2).value = product["productCode"]
            changed_refs.append(ws.cell(row, 2).coordinate)
        if product.get("insuranceCode") and current_insurance_code != product["insuranceCode"]:
            ws.cell(row, 6).value = product["insuranceCode"]
            changed_refs.append(ws.cell(row, 6).coordinate)
        if new_limit:
            if norm(ws.cell(row, 13).value) != new_limit:
                ws.cell(row, 13).value = new_limit
                changed_refs.append(ws.cell(row, 13).coordinate)

    highlight_changed(ws, changed_refs, YELLOW_FILL)

    # Rebuild support sheets from current product list.
    for name in ["변경히스토리", "지원Agent_수정필요항목", "지원Agent_1장요약"]:
        if name in wb.sheetnames:
            del wb[name]

    history = wb.create_sheet("변경히스토리")
    history_rows = [
        ["지원Agent 변경히스토리"],
        ["일시", "상품코드", "보험코드", "특약명", "변경유형", "변경내용", "비고", None, None, None, "메모"],
    ]
    support_rows = [
        ["지원Agent 수정필요항목"],
        ["상품코드", "보험코드", "특약명", "변경항목", "가입한도", "변경필요사유", "검토메모", "처리상태"],
    ]
    summary_rows = [["지원Agent 1장 요약"]]

    for product in products:
        change_text = format_limit_pair(product)
        if not change_text:
            continue
        change_info = classify_limit_change(product) or {"status": "룰검토필요", "reason": "변경 판단 필요"}
        status = change_info["status"]
        history_rows.append(
            [
                "2026-06-16",
                product.get("productCode") or "LP0000001",
                product.get("insuranceCode") or "",
                label_name(product),
                status,
                change_text or "변경 없음",
                "",
                None,
                None,
                None,
                review_note(product) or "변경 없음",
            ]
        )
        support_rows.append(
            [
                product.get("productCode") or "LP0000001",
                product.get("insuranceCode") or "",
                label_name(product),
                "가입한도금액RD",
                change_text,
                change_info["reason"],
                review_note(product) or "변경 없음",
                status,
            ]
        )
        summary_rows.append([])
        summary_rows.append(
            [
                f"[{status}] {label_name(product)} ({product.get('insuranceCode') or ''})\n"
                f"상품코드: {product.get('productCode') or ''}\n"
                f"변경항목: 가입한도금액RD\n"
                f"제안: {change_text or '변경 없음'}\n"
                f"메모: {review_note(product) or '변경 없음'}"
            ]
        )
        summary_rows.append([])

    for row in history_rows:
        history.append(row)
    support = wb.create_sheet("지원Agent_수정필요항목")
    for row in support_rows:
        support.append(row)
    summary = wb.create_sheet("지원Agent_1장요약")
    for row in summary_rows:
        summary.append(row)

    # History formatting
    history["A1"].fill = TITLE_FILL
    history["A1"].font = TITLE_FONT
    history["A1"].alignment = CENTER
    history.merge_cells("A1:K1")
    for cell in history[2]:
        cell.fill = BLUE_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BODY_BORDER
    for row in range(3, history.max_row + 1):
        for col in range(1, 12):
            cell = history.cell(row, col)
            cell.border = BODY_BORDER
            cell.alignment = TOP
            if col == 5:
                if cell.value == "자동반영":
                    cell.fill = GREEN_FILL
                elif cell.value == "룰검토필요":
                    cell.fill = ORANGE_FILL
    for col, width in {"A": 12, "B": 14, "C": 14, "D": 14, "E": 12, "F": 32, "G": 10, "H": 10, "I": 10, "J": 10, "K": 36}.items():
        history.column_dimensions[col].width = width

    # Support formatting
    support["A1"].fill = TITLE_FILL
    support["A1"].font = TITLE_FONT
    support["A1"].alignment = CENTER
    support.merge_cells("A1:H1")
    for cell in support[2]:
        cell.fill = BLUE_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BODY_BORDER
    for row in range(3, support.max_row + 1):
        for col in range(1, 9):
            cell = support.cell(row, col)
            cell.border = BODY_BORDER
            cell.alignment = TOP
            if col == 8:
                if cell.value == "자동반영":
                    cell.fill = GREEN_FILL
                elif cell.value == "룰검토필요":
                    cell.fill = ORANGE_FILL
    for col, width in {"A": 14, "B": 14, "C": 16, "D": 16, "E": 24, "F": 26, "G": 38, "H": 14}.items():
        support.column_dimensions[col].width = width

    # Summary formatting
    summary["A1"].fill = TITLE_FILL
    summary["A1"].font = TITLE_FONT
    summary["A1"].alignment = CENTER
    for row_idx in range(2, summary.max_row + 1):
        cell = summary.cell(row_idx, 1)
        if not norm(cell.value):
            continue
        cell.alignment = TOP
        cell.border = BODY_BORDER
        cell.fill = GREEN_FILL if "[자동반영]" in str(cell.value) else ORANGE_FILL
        cell.font = Font(bold=True, color="17202E")
        summary.row_dimensions[row_idx].height = 48
    for col in "ABCDEFGHIJ":
        summary.column_dimensions[col].width = 18

    wb.save(RD_OUT)


def build_guideline_summary(products):
    wb = load_workbook(GUIDELINE_PATH)
    if "지원Agent_수정필요항목" in wb.sheetnames:
        del wb["지원Agent_수정필요항목"]
    support = wb.create_sheet("지원Agent_수정필요항목")
    rows = [
        ["지원Agent 수정필요항목"],
        ["상품코드", "보험코드", "특약명", "변경항목", "가입한도", "변경필요사유", "검토메모", "처리상태"],
    ]
    for product in products:
        change_text = format_limit_pair(product)
        if not change_text:
            continue
        change_info = classify_limit_change(product) or {"status": "룰검토필요", "reason": "변경 판단 필요"}
        status = change_info["status"]
        rows.append(
            [
                product.get("productCode") or "",
                product.get("insuranceCode") or "",
                label_name(product),
                "가입한도금액RD",
                change_text,
                change_info["reason"],
                review_note(product) or "변경 없음",
                status,
            ]
        )
    for row in rows:
        support.append(row)
    support["A1"].fill = TITLE_FILL
    support["A1"].font = TITLE_FONT
    support["A1"].alignment = CENTER
    support.merge_cells(f"A1:H1")
    for cell in support[2]:
        cell.fill = BLUE_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BODY_BORDER
    for row in range(3, support.max_row + 1):
        for col in range(1, 9):
            cell = support.cell(row, col)
            cell.border = BODY_BORDER
            cell.alignment = TOP
            if col == 8:
                if cell.value == "자동반영":
                    cell.fill = GREEN_FILL
                elif cell.value == "룰검토필요":
                    cell.fill = ORANGE_FILL
    for col, width in {"A": 14, "B": 14, "C": 16, "D": 16, "E": 24, "F": 26, "G": 38, "H": 14}.items():
        support.column_dimensions[col].width = width
    wb.save(GUIDELINE_OUT)


def main():
    OUTPUTS.mkdir(parents=True, exist_ok=True)
    products = load_products()
    build_rd_workbook(products)
    build_guideline_summary(products)
    print(f"updated: {RD_OUT}")
    print(f"updated: {GUIDELINE_OUT}")


if __name__ == "__main__":
    main()
